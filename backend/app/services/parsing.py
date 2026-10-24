"""Resume file validation and text/metadata extraction.

Uploads are treated as hostile input:
  * the file type comes from the bytes (magic numbers), never from the name or Content-Type
  * DOCX archives are inspected before parsing (member count, uncompressed size) to stop zip bombs
  * PDF parsing is capped by page count and output length
"""
import io
import re
import unicodedata
import zipfile
from datetime import date

from ..errors import AppError

MAX_PDF_PAGES = 15
MAX_TEXT_CHARS = 60_000
MIN_TEXT_CHARS = 200
_DOCX_MAX_MEMBERS = 400
_DOCX_MAX_UNCOMPRESSED = 40 * 1024 * 1024


def detect_file_type(data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                infos = zf.infolist()
                if len(infos) > _DOCX_MAX_MEMBERS or sum(i.file_size for i in infos) > _DOCX_MAX_UNCOMPRESSED:
                    raise AppError(422, "This file is too complex to process", "invalid_file")
                names = {i.filename for i in infos}
                if "word/document.xml" in names and "[Content_Types].xml" in names:
                    if any(n.lower().startswith("word/vbaproject") for n in names):
                        raise AppError(415, "Macro-enabled documents are not accepted", "invalid_file")
                    return "docx"
        except zipfile.BadZipFile:
            pass
    raise AppError(415, "Upload a PDF or Word (.docx) file", "unsupported_file")


def safe_display_name(name: str | None, ext: str) -> str:
    """A filename that is safe to show and to put in a Content-Disposition header."""
    base = (name or "resume").replace("\\", "/").rsplit("/", 1)[-1]
    base = unicodedata.normalize("NFKD", base).encode("ascii", "ignore").decode()
    base = re.sub(r"[^A-Za-z0-9._ -]", "", base).strip(" .") or "resume"
    stem = base.rsplit(".", 1)[0][:100] or "resume"
    return f"{stem}.{ext}"


def _clean(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).replace("\x00", "")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


def extract_text(data: bytes, file_type: str) -> str:
    if file_type == "pdf":
        import pymupdf
        try:
            doc = pymupdf.open(stream=data, filetype="pdf")
        except Exception:
            raise AppError(422, "This PDF could not be opened", "invalid_file")
        try:
            if doc.needs_pass:
                raise AppError(422, "Password-protected PDFs can't be processed", "invalid_file")
            parts, total = [], 0
            for i, page in enumerate(doc):
                if i >= MAX_PDF_PAGES or total > MAX_TEXT_CHARS:
                    break
                chunk = page.get_text("text")
                parts.append(chunk)
                total += len(chunk)
            text = "\n".join(parts)
        finally:
            doc.close()
    else:
        import docx
        try:
            d = docx.Document(io.BytesIO(data))
        except Exception:
            raise AppError(422, "This Word document could not be opened", "invalid_file")
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    text = _clean(text)
    if len(text) < MIN_TEXT_CHARS:
        raise AppError(422, "We couldn't find readable text in this file. If it's a scanned image, "
                            "export a text-based PDF or upload a .docx instead.", "no_text")
    return text


# --- Skill matching ---------------------------------------------------------------------------------

SKILL_ALIASES = {
    "js": "JavaScript", "javascript": "JavaScript", "ts": "TypeScript", "golang": "Go", "py": "Python",
    "python3": "Python", "reactjs": "React", "react.js": "React", "nodejs": "Node.js", "node": "Node.js",
    "vue": "Vue.js", "vuejs": "Vue.js", "nextjs": "Next.js", "express": "Express.js", "k8s": "Kubernetes",
    "postgres": "PostgreSQL", "psql": "PostgreSQL", "mongo": "MongoDB", "gcp": "Google Cloud",
    "amazon web services": "AWS", "s3": "Amazon S3", "ec2": "Amazon EC2", "lambda": "AWS Lambda",
    "ml": "Machine Learning", "dl": "Deep Learning", "natural language processing": "NLP",
    "cv": "Computer Vision", "sklearn": "scikit-learn", "scikit learn": "scikit-learn", "tf": "TensorFlow",
    "huggingface": "Hugging Face", "transformers": "Hugging Face", "llm": "LLMs", "large language models": "LLMs",
    "retrieval augmented generation": "RAG", "spark": "Apache Spark", "pyspark": "Apache Spark",
    "ci cd": "CI/CD", "cicd": "CI/CD", "oop": "Object-Oriented Programming", "dsa": "Data Structures",
    "rest": "REST APIs", "rest api": "REST APIs", "restful apis": "REST APIs", "powerbi": "Power BI",
    "ms excel": "Excel", "microsoft excel": "Excel", "tailwind": "Tailwind CSS", "html5": "HTML", "css3": "CSS",
    "springboot": "Spring Boot", "spring": "Spring Boot", "dotnet": "ASP.NET", ".net": "ASP.NET",
    "github actions": "GitHub Actions", "mssql": "SQL Server", "dynamo db": "DynamoDB", "elastic search": "Elasticsearch",
    "dockerized": "Docker", "dockerised": "Docker", "kubernetes cluster": "Kubernetes", "postgre": "PostgreSQL",
}

# Aliases safe to look for in free text (the short ones above, like "js" or "ml", are too ambiguous).
_TEXT_ALIASES = {k: v for k, v in SKILL_ALIASES.items()
                 if len(k) >= 4 and k not in {"node", "rest", "spring", "express", "lambda", "transformers", "vue", "tf"}} | {"k8s": "Kubernetes"}

# Names that are also everyday words/letters: never scanned from raw text, only accepted when the
# structured extraction lists them explicitly.
_AMBIGUOUS = {"go", "r", "c", "rest", "git", "communication", "leadership", "problem solving", "statistics",
              "networking", "excel", "agile", "android", "ios", "swift", "dart", "spring", "express"}


def normalize_skill(raw: str, dictionary: dict[str, int]) -> int | None:
    """Map a free-text skill to a dictionary skill id. `dictionary` is {lower(name): skill_id}."""
    key = re.sub(r"\s+", " ", raw.strip().lower())
    if not key or len(key) > 60:
        return None
    if key in dictionary:
        return dictionary[key]
    alias = SKILL_ALIASES.get(key) or SKILL_ALIASES.get(key.replace("-", " "))
    return dictionary.get(alias.lower()) if alias else None


def _mentions(term: str, lower_text: str) -> bool:
    return re.search(r"(?<![a-z0-9+#.])" + re.escape(term) + r"(?![a-z0-9+#])", lower_text) is not None


def scan_text_for_skills(text: str, dictionary: dict[str, int]) -> set[int]:
    """Dictionary skills mentioned in free text, including common aliases ("Postgres", "k8s", "Python3")."""
    found: set[int] = set()
    lower = text.lower()
    for name, skill_id in dictionary.items():
        if name in _AMBIGUOUS or len(name) < 2:
            continue
        if _mentions(name, lower):
            found.add(skill_id)
    for alias, canonical in _TEXT_ALIASES.items():
        sid = dictionary.get(canonical.lower())
        if sid and sid not in found and _mentions(alias, lower):
            found.add(sid)
    return found


_SKILL_LIST_LINE = re.compile(r"^\s*(technical\s+)?(skills|technologies|tech stack|tools)\s*[:|-]", re.I)


def without_skill_lists(text: str) -> str:
    """The resume minus its 'Skills:' style lines: what's left is where skills are *used*."""
    return "\n".join(line for line in text.splitlines() if not _SKILL_LIST_LINE.match(line))


# --- Rule-based fallback extraction (used when the LLM is unavailable) ------------------------------------

_YEARS_RE = re.compile(r"(\d{1,2}(?:\.\d)?)\s*\+?\s*(?:years|yrs)(?:\s+of)?\s+(?:experience|exp)", re.I)
_EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_RE = re.compile(r"\+?\d[\d ()-]{8,16}\d")
_DEGREE_RE = re.compile(r"\b(B\.?\s?Tech|B\.?\s?E\.?|B\.?\s?Sc|BCA|MCA|M\.?\s?Tech|M\.?\s?Sc|MBA|Ph\.?D|Bachelor[^\n,]{0,60}|Master[^\n,]{0,60})", re.I)


_RANGE_RE = re.compile(r"\b((?:19|20)\d{2})\s*(?:-|\u2013|\u2014|to)\s*((?:19|20)\d{2}|present|current|now|date)\b", re.I)
_SINCE_RE = re.compile(r"\bsince\s+((?:19|20)\d{2})\b", re.I)
# Whole words only: the earlier substring version matched "nit" inside "monitoring" and "unit tests",
# and so silently dropped real work history from the experience total (found by eval/run_eval.py).
_EDU_WORDS = re.compile(
    r"(?<![a-z])(b\.?\s?tech|b\.e\.?|bca|mca|m\.?\s?tech|m\.?\s?sc|b\.?\s?sc|b\.?\s?com|bba|mba|ph\.?\s?d|"
    r"universit(y|ies)|college|institute|school|degree|bachelor'?s?|masters|master of|iit|nit|iisc|bits pilani)(?![a-z])", re.I)
# Career breaks, sabbaticals etc. are not work experience (also found by the evaluation).
_BREAK_WORDS = re.compile(r"(?<![a-z])(career break|sabbatical|employment gap|gap year|break from work)(?![a-z])", re.I)


def _years_from_date_ranges(text: str) -> float | None:
    """Estimate experience from 'YYYY - YYYY/Present' ranges on non-education lines, merging overlaps."""
    this_year = date.today().year
    spans = []
    for line in text.splitlines():
        if _EDU_WORDS.search(line) or _BREAK_WORDS.search(line):
            continue
        for m in _RANGE_RE.finditer(line):
            start = int(m.group(1))
            end = this_year if not m.group(2).isdigit() else int(m.group(2))
            if 1970 <= start <= end <= this_year:
                spans.append((start, end))
        if not _RANGE_RE.search(line):
            for m in _SINCE_RE.finditer(line):            # "at Payly since 2020"
                start = int(m.group(1))
                if 1970 <= start <= this_year:
                    spans.append((start, this_year))
    if not spans:
        return None
    spans.sort()
    total, (cur_s, cur_e) = 0, spans[0]
    for s_, e_ in spans[1:]:
        if s_ <= cur_e:
            cur_e = max(cur_e, e_)
        else:
            total += cur_e - cur_s
            cur_s, cur_e = s_, e_
    total += cur_e - cur_s
    return float(min(total, 50)) or None


def rule_based_extract(text: str) -> dict:
    years = [float(m.group(1)) for m in _YEARS_RE.finditer(text)]
    if not years and (from_ranges := _years_from_date_ranges(text)):
        years = [from_ranges]
    first_line = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")
    return {
        "name": first_line[:80] if len(first_line.split()) <= 5 else None,
        "email": (_EMAIL_RE.search(text) or [None])[0],
        "phone": (_PHONE_RE.search(text) or [None])[0],
        "headline": None,
        "location": None,
        "total_experience_years": max(years) if years else None,
        "skills": [],
        "education": [{"degree": m.group(0).strip()[:120], "institution": None, "field_of_study": None,
                       "start_year": None, "end_year": None} for m in list(_DEGREE_RE.finditer(text))[:3]],
        "experience": [],
        "projects": [],
        "certifications": [],
        "summary": None,
        "extracted_by": "rules",
    }


_URL_RE = re.compile(r"(https?://\S+|www\.\S+|\b[\w-]+\.(?:com|in|io|dev|me|org)/\S*)", re.I)


def embedding_text(text: str, redact: list[str] | None = None) -> str:
    """Resume text for the semantic embedding, with everything that should not affect job fit removed:
    names, contact details, links, location, education lines (institution prestige) and career-break
    lines. The counterfactual tests in eval/ check that changing these leaves the score unchanged."""
    lines = [ln for ln in text.splitlines() if not _EDU_WORDS.search(ln) and not _BREAK_WORDS.search(ln)]
    out = "\n".join(lines)
    out = _EMAIL_RE.sub(" ", out)
    out = _URL_RE.sub(" ", out)
    out = _PHONE_RE.sub(" ", out)
    for term in sorted({t.strip() for t in (redact or []) if t and len(t.strip()) >= 2}, key=len, reverse=True):
        out = re.sub(r"(?<![\w])" + re.escape(term) + r"(?![\w])", " ", out, flags=re.I)
    return re.sub(r"[ \t]+", " ", out).strip()


def parse_month(value) -> date | None:
    """Parse 'YYYY-MM', 'YYYY' or 'Present'-like values from the extractor into a date."""
    if not value or not isinstance(value, str):
        return None
    m = re.match(r"^\s*(\d{4})(?:-(\d{1,2}))?", value)
    if not m:
        return None
    year, month = int(m.group(1)), int(m.group(2) or 1)
    if not (1950 <= year <= date.today().year + 1 and 1 <= month <= 12):
        return None
    return date(year, month, 1)
